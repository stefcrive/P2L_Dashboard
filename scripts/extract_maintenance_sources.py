from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
SOURCE_ROOT = REPOSITORY_ROOT / "Mantainance"
OUTPUT_ROOT = REPOSITORY_ROOT / "tmp" / "maintenance_cleaning"
RAW_OUTPUT = OUTPUT_ROOT / "raw_sources.json"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def trim_rows(rows: list[list[Any]]) -> list[list[Any]]:
    trimmed: list[list[Any]] = []
    for row in rows:
        last = -1
        for index, value in enumerate(row):
            if value not in (None, ""):
                last = index
        if last >= 0:
            trimmed.append(row[: last + 1])
    return trimmed


def extract_workbook(path: Path) -> dict[str, Any]:
    formula_book = load_workbook(path, data_only=False, read_only=False)
    value_book = load_workbook(path, data_only=True, read_only=False)
    sheets: list[dict[str, Any]] = []

    for formula_sheet in formula_book.worksheets:
        value_sheet = value_book[formula_sheet.title]
        rows: list[list[Any]] = []
        formulas: dict[str, str] = {}
        for row_index in range(1, formula_sheet.max_row + 1):
            values: list[Any] = []
            for column_index in range(1, formula_sheet.max_column + 1):
                formula_cell = formula_sheet.cell(row_index, column_index)
                value_cell = value_sheet.cell(row_index, column_index)
                if isinstance(formula_cell.value, str) and formula_cell.value.startswith("="):
                    formulas[formula_cell.coordinate] = formula_cell.value
                    value = value_cell.value if value_cell.value is not None else formula_cell.value
                else:
                    value = formula_cell.value
                values.append(clean_value(value))
            rows.append(values)

        sheets.append(
            {
                "name": formula_sheet.title,
                "max_row": formula_sheet.max_row,
                "max_column": formula_sheet.max_column,
                "hidden": formula_sheet.sheet_state != "visible",
                "merged_ranges": [str(item) for item in formula_sheet.merged_cells.ranges],
                "rows": trim_rows(rows),
                "formulas": formulas,
            }
        )

    return {"kind": "spreadsheet", "sheets": sheets}


def read_csv_text(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace"), "latin-1-replace"


def extract_csv(path: Path) -> dict[str, Any]:
    text, encoding = read_csv_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";"
    rows = [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    return {
        "kind": "csv",
        "encoding": encoding,
        "delimiter": delimiter,
        "rows": trim_rows(rows),
    }


def extract_pdf(path: Path) -> dict[str, Any]:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append({"page": index, "text": text})
    return {
        "kind": "pdf",
        "page_count": len(reader.pages),
        "pages": pages,
        "text_characters": sum(len(page["text"]) for page in pages),
    }


def main() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    records: list[dict[str, Any]] = []
    errors: list[dict[str, str]] = []

    for path in sorted(SOURCE_ROOT.rglob("*")):
        if not path.is_file():
            continue
        relative_path = path.relative_to(SOURCE_ROOT).as_posix()
        extension = path.suffix.lower()
        record: dict[str, Any] = {
            "relative_path": relative_path,
            "extension": extension,
            "size_bytes": path.stat().st_size,
            "sha256": file_sha256(path),
        }
        try:
            if extension in {".xlsx", ".xlsm"}:
                record.update(extract_workbook(path))
            elif extension == ".csv":
                record.update(extract_csv(path))
            elif extension == ".pdf":
                record.update(extract_pdf(path))
            else:
                record["kind"] = "unsupported"
            records.append(record)
        except Exception as exc:  # pragma: no cover - audit capture
            errors.append({"relative_path": relative_path, "error": str(exc)})

    payload = {
        "source_root": str(SOURCE_ROOT),
        "file_count": len(records),
        "error_count": len(errors),
        "errors": errors,
        "files": records,
    }
    RAW_OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output": str(RAW_OUTPUT),
        "file_count": len(records),
        "error_count": len(errors),
        "kinds": {
            kind: sum(1 for item in records if item.get("kind") == kind)
            for kind in sorted({str(item.get("kind")) for item in records})
        },
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
